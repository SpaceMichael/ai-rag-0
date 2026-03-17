"""
app/ingestion/loader.py — 多類型文件載入器
支援：PDF, Word, CSV, Excel, TXT, HTML, PowerPoint
"""
import os
from pathlib import Path
from dataclasses import dataclass
from loguru import logger
import pandas as pd
from bs4 import BeautifulSoup


@dataclass
class RawDocument:
    """載入後的原始文件"""
    content: str
    source: str
    metadata: dict


class DocumentLoader:
    """萬能文件載入器，根據副檔名自動選擇解析方式"""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx", ".html", ".htm", ".pptx"}

    def load(self, file_path: str | Path) -> RawDocument:
        """載入單個文件"""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"不支援的文件類型: {ext}")

        logger.info(f"📄 載入文件: {path.name} ({ext})")

        loaders = {
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
            ".txt": self._load_txt,
            ".csv": self._load_csv,
            ".xlsx": self._load_excel,
            ".html": self._load_html,
            ".htm": self._load_html,
            ".pptx": self._load_pptx,
        }
        return loaders[ext](path)

    def load_directory(self, dir_path: str | Path) -> list[RawDocument]:
        """批量載入目錄下所有支援的文件"""
        path = Path(dir_path)
        docs = []
        for file in path.rglob("*"):
            if file.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                try:
                    docs.append(self.load(file))
                except Exception as e:
                    logger.warning(f"⚠️ 跳過文件 {file.name}: {e}")
        logger.info(f"📚 共載入 {len(docs)} 個文件")
        return docs

    def _load_pdf(self, path: Path) -> RawDocument:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return RawDocument(content=text, source=path.name, metadata={"pages": len(reader.pages), "type": "pdf"})

    def _load_docx(self, path: Path) -> RawDocument:
        from docx import Document
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return RawDocument(content=text, source=path.name, metadata={"type": "docx"})

    def _load_txt(self, path: Path) -> RawDocument:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return RawDocument(content=text, source=path.name, metadata={"type": "txt"})

    def _load_csv(self, path: Path) -> RawDocument:
        df = pd.read_csv(path)
        text = df.to_string(index=False)
        return RawDocument(content=text, source=path.name, metadata={"rows": len(df), "type": "csv"})

    def _load_excel(self, path: Path) -> RawDocument:
        df = pd.read_excel(path, sheet_name=None)
        parts = []
        for sheet_name, sheet_df in df.items():
            parts.append(f"[試算表: {sheet_name}]\n{sheet_df.to_string(index=False)}")
        text = "\n\n".join(parts)
        return RawDocument(content=text, source=path.name, metadata={"sheets": list(df.keys()), "type": "xlsx"})

    def _load_html(self, path: Path) -> RawDocument:
        html = path.read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(html, "lxml")
        for tag in soup(["script", "style"]):
            tag.decompose()
        text = soup.get_text(separator="\n", strip=True)
        return RawDocument(content=text, source=path.name, metadata={"type": "html"})

    def _load_pptx(self, path: Path) -> RawDocument:
        from pptx import Presentation
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            if slide_text:
                parts.append(f"[第{i}頁]\n" + "\n".join(slide_text))
        text = "\n\n".join(parts)
        return RawDocument(content=text, source=path.name, metadata={"slides": len(prs.slides), "type": "pptx"})
